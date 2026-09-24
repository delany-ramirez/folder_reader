"""Escritura del inventario a un archivo .xlsx."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FORMATO_FECHA = "yyyy-mm-dd hh:mm:ss"

# Excel admite 1.048.576 filas contando el encabezado.
MAX_FILAS_DATOS = 1_048_575

COLUMNAS: tuple[tuple[str, int], ...] = (
    ("Ruta completa", 70),
    ("Ruta relativa", 45),
    ("Carpeta contenedora", 45),
    ("Carpeta nivel 1", 18),
    ("Nombre del archivo", 40),
    ("Extension", 12),
    ("Tamano (bytes)", 16),
    ("Tamano (MB)", 13),
    ("Fecha de creacion", 21),
    ("Fecha de ultima modificacion", 28),
    ("Guardado el (documento)", 24),
    ("Numero de paginas", 18),
    ("Estado", 40),
)


@dataclass(slots=True)
class Fila:
    """Una fila del inventario, ya lista para escribirse."""

    ruta: str
    ruta_relativa: str
    carpeta: str
    carpeta_nivel1: str
    nombre: str
    extension: str
    tamano_bytes: int
    creado: datetime | None
    modificado: datetime | None
    guardado: datetime | None
    paginas: int | None
    estado: str

    def valores(self) -> list[object]:
        return [
            _limpiar(self.ruta),
            _limpiar(self.ruta_relativa),
            _limpiar(self.carpeta),
            _limpiar(self.carpeta_nivel1),
            _limpiar(self.nombre),
            _limpiar(self.extension),
            self.tamano_bytes,
            round(self.tamano_bytes / 1_048_576, 2),
            self.creado,
            self.modificado,
            self.guardado,
            self.paginas,
            _limpiar(self.estado),
        ]


def _limpiar(texto: str) -> str:
    """Quita caracteres de control que openpyxl rechaza al guardar."""
    return ILLEGAL_CHARACTERS_RE.sub("", texto)


def escribir_excel(filas: list[Fila], destino: Path) -> int:
    """Escribe las filas en `destino` y devuelve cuantas se guardaron."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Inventario"

    hoja.append([titulo for titulo, _ in COLUMNAS])

    relleno = PatternFill("solid", fgColor="1F4E78")
    negrita = Font(bold=True, color="FFFFFF")
    for celda in hoja[1]:
        celda.fill = relleno
        celda.font = negrita
        celda.alignment = Alignment(vertical="center", wrap_text=True)

    escritas = filas[:MAX_FILAS_DATOS]
    for fila in escritas:
        hoja.append(fila.valores())

    for indice, (_, ancho) in enumerate(COLUMNAS, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho

    # Las tres columnas de fecha se guardan como datetime para que Excel las
    # ordene y filtre como fechas reales, no como texto.
    primera_fecha = 1 + [t for t, _ in COLUMNAS].index("Fecha de creacion")
    for fila_excel in hoja.iter_rows(
        min_row=2, min_col=primera_fecha, max_col=primera_fecha + 2
    ):
        for celda in fila_excel:
            celda.number_format = FORMATO_FECHA

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    destino.parent.mkdir(parents=True, exist_ok=True)
    libro.save(destino)
    return len(escritas)
