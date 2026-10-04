"""Escritura de la hoja de control documental a partir de una plantilla .xlsx.

La plantilla manda en todo lo visual: encabezado con logo, anchos, fuentes,
bordes, fila TOTAL y pie de firmas. De ella se toman dos filas modelo (la de
carpeta, con el nombre en negrilla, y la de documento) y se replican tantas
veces como haga falta; luego se reponen la fila TOTAL y el pie debajo.
"""

from __future__ import annotations

import re
import unicodedata
from copy import copy
from dataclasses import dataclass, field
from datetime import date
from importlib import resources
from importlib.resources.abc import Traversable
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

FORMATO_FECHA = "dd/mm/yyyy"

# Titulos de la fila de encabezado que identifican cada columna de la plantilla.
ENCABEZADOS = {
    "item": "ITEM",
    "fecha_documento": "FECHA DOCUMENTO",
    "nombre": "TIPO DOCUMENTAL",
    "folios": "CANTIDAD DE FOLIOS",
    "fecha_registro": "FECHA REGISTRO",
}
TEXTO_TOTAL = "TOTAL"


class ErrorPlantilla(Exception):
    """La plantilla no tiene la estructura esperada."""


@dataclass(slots=True)
class Documento:
    """Un archivo de la hoja de control."""

    nombre: str
    fecha: date | None  # de guardado del documento; None si no la registra
    folios: int | None


@dataclass(slots=True)
class Grupo:
    """Documentos de una misma carpeta. `carpeta` vacia = sin fila de carpeta."""

    carpeta: str
    documentos: list[Documento] = field(default_factory=list)


@dataclass(slots=True)
class _Bloque:
    """Filas copiadas de la plantilla, con coordenadas relativas a la primera."""

    estilos: dict[tuple[int, int], object]
    valores: dict[tuple[int, int], object]
    altos: list[float | None]
    combinadas: list[tuple[int, int, int]]  # (fila relativa, col. inicial, col. final)


def plantilla_predeterminada() -> Traversable:
    return resources.files("folder_reader") / "plantillas" / "hoja_control.xlsx"


def clave_natural(texto: str) -> tuple[list[object], str]:
    """Ordena los numeros por su valor: 1, 2, 10, 11 y no 1, 10, 11, 2.

    Es el orden del Explorador de Windows: "2. RP" antes que "10. CEDULA" y
    "2.9" antes que "2.10". NFKC convierte los digitos de ancho completo (１２)
    en digitos normales; el texto original desempata nombres que solo difieren
    en mayusculas o ceros a la izquierda ("01" y "1").
    """
    normal = unicodedata.normalize("NFKC", texto).strip()
    partes = re.split(r"([0-9]+)", normal)
    return [int(t) if t.isdigit() else t.casefold() for t in partes], texto


def _limpiar(texto: str) -> str:
    """Quita caracteres de control que openpyxl rechaza al guardar."""
    return ILLEGAL_CHARACTERS_RE.sub("", texto)


def _normalizar(valor: object) -> str:
    """Texto en mayusculas, sin tildes ni espacios sobrantes, para comparar titulos."""
    if not isinstance(valor, str):
        return ""
    sin_tildes = unicodedata.normalize("NFKD", valor).encode("ascii", "ignore").decode()
    return " ".join(sin_tildes.split()).upper()


def _es_suma(valor: object) -> bool:
    return isinstance(valor, str) and valor.upper().startswith("=SUM(")


def _capturar(
    hoja: Worksheet, primera: int, ultima: int, max_col: int, *, con_valores: bool
) -> _Bloque:
    estilos: dict[tuple[int, int], object] = {}
    valores: dict[tuple[int, int], object] = {}
    for relativa, fila in enumerate(range(primera, ultima + 1)):
        for col in range(1, max_col + 1):
            celda = hoja.cell(fila, col)
            estilos[relativa, col] = copy(celda._style)
            if con_valores and celda.value is not None:
                valores[relativa, col] = celda.value
    combinadas = [
        (rango.min_row - primera, rango.min_col, rango.max_col)
        for rango in hoja.merged_cells.ranges
        if primera <= rango.min_row == rango.max_row <= ultima
    ]
    altos = [hoja.row_dimensions[f].height for f in range(primera, ultima + 1)]
    return _Bloque(estilos, valores, altos, combinadas)


def _pegar(hoja: Worksheet, bloque: _Bloque, fila: int) -> None:
    """Reproduce `bloque` a partir de `fila`: combinaciones, estilos, valores y altos."""
    for relativa, col_ini, col_fin in bloque.combinadas:
        hoja.merge_cells(
            start_row=fila + relativa,
            start_column=col_ini,
            end_row=fila + relativa,
            end_column=col_fin,
        )
    # Se combina antes de dar estilo: las celdas cubiertas pasan a ser MergedCell,
    # que tambien admiten estilo, y asi conservan los bordes de la plantilla.
    for (relativa, col), estilo in bloque.estilos.items():
        hoja.cell(fila + relativa, col)._style = copy(estilo)
    for (relativa, col), valor in bloque.valores.items():
        hoja.cell(fila + relativa, col).value = valor
    for relativa, alto in enumerate(bloque.altos):
        hoja.row_dimensions[fila + relativa].height = alto


class _Plantilla:
    """Ubica en la hoja las piezas que se van a replicar."""

    def __init__(self, hoja: Worksheet) -> None:
        self.hoja = hoja
        self.max_col = hoja.max_column
        self.fila_encabezado = self._buscar_encabezado()
        self.columnas = self._mapear_columnas()

        fila_total = next(
            (
                f
                for f in range(self.fila_encabezado + 1, hoja.max_row + 1)
                if any(
                    _normalizar(hoja.cell(f, c).value) == TEXTO_TOTAL
                    for c in range(1, self.max_col + 1)
                )
            ),
            None,
        )
        if fila_total is None:
            raise ErrorPlantilla(f'La plantilla no tiene una fila "{TEXTO_TOTAL}".')

        datos = range(self.fila_encabezado + 1, fila_total)
        if not datos:
            raise ErrorPlantilla(
                "La plantilla necesita al menos una fila de documento entre el "
                f'encabezado y la fila "{TEXTO_TOTAL}".'
            )
        # Modelo de carpeta: primera fila con el nombre en negrilla. Modelo de
        # documento: primera fila con el nombre sin negrilla.
        col_nombre = self.columnas["nombre"]
        negrilla = [f for f in datos if hoja.cell(f, col_nombre).font.b]
        normal = [f for f in datos if not hoja.cell(f, col_nombre).font.b]
        fila_documento = normal[0] if normal else datos[0]
        fila_carpeta = negrilla[0] if negrilla else fila_documento

        # De las filas modelo solo interesa el formato; el total y el pie se
        # copian con sus textos (TOTAL, firmas).
        self.carpeta = _capturar(
            hoja, fila_carpeta, fila_carpeta, self.max_col, con_valores=False
        )
        self.documento = _capturar(
            hoja, fila_documento, fila_documento, self.max_col, con_valores=False
        )
        self.pie = _capturar(hoja, fila_total, hoja.max_row, self.max_col, con_valores=True)

    def _buscar_encabezado(self) -> int:
        objetivo = ENCABEZADOS["item"]
        for fila in self.hoja.iter_rows(max_col=self.max_col):
            for celda in fila:
                if _normalizar(celda.value) == objetivo:
                    return celda.row
        raise ErrorPlantilla(f'La plantilla no tiene una columna "{objetivo}".')

    def _mapear_columnas(self) -> dict[str, int]:
        titulos = {
            _normalizar(self.hoja.cell(self.fila_encabezado, c).value): c
            for c in range(1, self.max_col + 1)
        }
        faltan = [t for t in ENCABEZADOS.values() if t not in titulos]
        if faltan:
            raise ErrorPlantilla(
                "Faltan columnas en el encabezado de la plantilla: " + ", ".join(faltan)
            )
        return {clave: titulos[titulo] for clave, titulo in ENCABEZADOS.items()}

    def vaciar(self) -> None:
        """Borra todo lo que hay bajo el encabezado: filas de ejemplo, total y pie."""
        hoja = self.hoja
        limite = self.fila_encabezado
        for rango in list(hoja.merged_cells.ranges):
            if rango.min_row > limite:
                hoja.merged_cells.remove(rango)
        for clave in [k for k in hoja._cells if k[0] > limite]:
            del hoja._cells[clave]
        for fila in [f for f in hoja.row_dimensions if f > limite]:
            del hoja.row_dimensions[fila]


def escribir_hoja_control(
    grupos: list[Grupo], destino: Path, plantilla: Path | None = None
) -> int:
    """Vuelca los grupos en una copia de la plantilla. Devuelve los documentos escritos.

    Las filas de carpeta y de documento llevan numero de ITEM consecutivo, como
    en la plantilla. Las dos fechas de cada documento son la de guardado, y
    quedan en blanco si el documento no la registra.
    """
    origen = plantilla if plantilla is not None else plantilla_predeterminada()
    with origen.open("rb") as fh:
        libro = load_workbook(fh)
    hoja = libro.active
    modelo = _Plantilla(hoja)
    columnas = modelo.columnas
    modelo.vaciar()

    primera_dato = fila = modelo.fila_encabezado + 1
    item = 1
    escritos = 0

    for grupo in grupos:
        if grupo.carpeta:
            _pegar(hoja, modelo.carpeta, fila)
            hoja.cell(fila, columnas["item"]).value = item
            hoja.cell(fila, columnas["nombre"]).value = _limpiar(grupo.carpeta)
            fila += 1
            item += 1

        for documento in grupo.documentos:
            _pegar(hoja, modelo.documento, fila)
            hoja.cell(fila, columnas["item"]).value = item
            hoja.cell(fila, columnas["nombre"]).value = _limpiar(documento.nombre)
            hoja.cell(fila, columnas["folios"]).value = documento.folios
            for clave in ("fecha_documento", "fecha_registro"):
                celda = hoja.cell(fila, columnas[clave])
                celda.value = documento.fecha
                celda.number_format = FORMATO_FECHA
            fila += 1
            item += 1
            escritos += 1

    fila_total = fila
    _pegar(hoja, modelo.pie, fila_total)
    # TOTAL no es un item; un numero ahi es un resto de la hoja de la que salio la plantilla.
    hoja.cell(fila_total, columnas["item"]).value = None

    # La suma de la plantilla apunta a sus filas de ejemplo: se reescribe sobre
    # el rango real. Sin documentos queda en cero.
    celdas_total = [hoja.cell(fila_total, c) for c in range(1, modelo.max_col + 1)]
    sumas = [c for c in celdas_total if _es_suma(c.value)] or [
        hoja.cell(fila_total, columnas["folios"])
    ]
    for celda in sumas:
        letra = get_column_letter(celda.column)
        celda.value = (
            f"=SUM({letra}{primera_dato}:{letra}{fila_total - 1})"
            if fila_total > primera_dato
            else 0
        )

    destino.parent.mkdir(parents=True, exist_ok=True)
    libro.save(destino)
    return escritos
