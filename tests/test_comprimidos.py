from __future__ import annotations

import os
import subprocess
import zipfile
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from folder_reader import cli
from folder_reader.comprimidos import (
    _RUTAS_WINRAR,
    _buscar,
    comprimidos_en_raiz,
    descomprimir_en_raiz,
)

from conftest import crear_pdf


def crear_zip(destino: Path, miembros: dict[str, bytes], fecha=(2020, 5, 6, 7, 8, 10)) -> Path:
    with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as zf:
        for nombre, contenido in miembros.items():
            zf.writestr(zipfile.ZipInfo(nombre, date_time=fecha), contenido)
    return destino


def test_descomprime_solo_los_de_la_raiz(arbol: Path) -> None:
    crear_zip(arbol / "entregas.zip", {"a.txt": b"a", "sub/b.txt": b"b"})
    crear_zip(arbol / "informes" / "anidado.zip", {"c.txt": b"c"})

    resultado = descomprimir_en_raiz(arbol)

    assert resultado.descomprimidos == ["entregas.zip"]
    assert (arbol / "entregas" / "a.txt").read_bytes() == b"a"
    assert (arbol / "entregas" / "sub" / "b.txt").is_file()
    assert not (arbol / "informes" / "anidado").exists()
    # El comprimido original se deja donde estaba.
    assert (arbol / "entregas.zip").is_file()


def test_no_pisa_una_carpeta_existente(arbol: Path) -> None:
    crear_zip(arbol / "informes.zip", {"nuevo.txt": b"x"})

    resultado = descomprimir_en_raiz(arbol)

    assert resultado.ya_existian == ["informes.zip"]
    assert not (arbol / "informes" / "nuevo.txt").exists()


def test_segunda_corrida_no_vuelve_a_extraer(arbol: Path) -> None:
    crear_zip(arbol / "entregas.zip", {"a.txt": b"a"})
    descomprimir_en_raiz(arbol)
    assert descomprimir_en_raiz(arbol).ya_existian == ["entregas.zip"]


def test_aplana_la_carpeta_con_el_mismo_nombre(arbol: Path) -> None:
    crear_zip(arbol / "Proyecto.zip", {"Proyecto/doc.txt": b"d", "Proyecto/x/y.txt": b"y"})

    descomprimir_en_raiz(arbol)

    assert (arbol / "Proyecto" / "doc.txt").is_file()
    assert (arbol / "Proyecto" / "x" / "y.txt").is_file()
    assert not (arbol / "Proyecto" / "Proyecto").exists()


def test_conserva_la_fecha_de_modificacion(arbol: Path) -> None:
    crear_zip(arbol / "viejo.zip", {"a.txt": b"a"}, fecha=(2020, 5, 6, 7, 8, 10))
    descomprimir_en_raiz(arbol)

    modificado = datetime.fromtimestamp((arbol / "viejo" / "a.txt").stat().st_mtime)
    assert modificado == datetime(2020, 5, 6, 7, 8, 10)


def test_comprimido_danado_no_interrumpe(arbol: Path) -> None:
    (arbol / "roto.zip").write_bytes(b"esto no es un zip")
    crear_zip(arbol / "bueno.zip", {"a.txt": b"a"})

    resultado = descomprimir_en_raiz(arbol)

    assert resultado.descomprimidos == ["bueno.zip"]
    assert [nombre for nombre, _ in resultado.fallidos] == ["roto.zip"]
    # No quedan carpetas a medias.
    assert not (arbol / "roto").exists()
    assert not any(p.name.endswith("~") for p in arbol.iterdir())


def test_zip_con_contrasena_se_reporta(arbol: Path) -> None:
    crear_zip(arbol / "secreto.zip", {"a.txt": b"a"})
    # Se marca el bit de cifrado en el directorio central y en la cabecera local.
    crudo = bytearray((arbol / "secreto.zip").read_bytes())
    for firma, desplazamiento in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
        posicion = crudo.find(firma)
        crudo[posicion + desplazamiento] |= 0x1
    (arbol / "secreto.zip").write_bytes(bytes(crudo))

    resultado = descomprimir_en_raiz(arbol)

    assert resultado.fallidos == [("secreto.zip", "Protegido con contrasena")]


def test_nombres_con_tildes_sin_marca_utf8(arbol: Path) -> None:
    """El Explorador guarda los nombres en la pagina OEM, sin la marca UTF-8."""
    from folder_reader.comprimidos import _codificacion_oem

    try:
        oem = "Año.txt".encode(_codificacion_oem())
    except UnicodeEncodeError:
        pytest.skip("la pagina OEM de este equipo no representa la enie")

    # Un nombre ASCII del mismo largo no lleva la marca UTF-8; luego se cambian
    # sus bytes por los del nombre real en la pagina OEM.
    provisional = b"A" + b"_" * (len(oem) - 6) + b"o.txt"
    destino = crear_zip(arbol / "oem.zip", {provisional.decode(): b"x"})
    destino.write_bytes(destino.read_bytes().replace(provisional, oem))

    descomprimir_en_raiz(arbol)

    assert (arbol / "oem" / "Año.txt").is_file()


def test_ignora_comprimidos_ocultos(arbol: Path) -> None:
    if os.name != "nt":
        pytest.skip("los atributos de Windows no existen en POSIX")
    crear_zip(arbol / "oculto.zip", {"a.txt": b"a"})
    subprocess.run(["attrib", "+H", str(arbol / "oculto.zip")], check=True, shell=True)

    assert comprimidos_en_raiz(arbol) == []
    assert [p.name for p in comprimidos_en_raiz(arbol, incluir_ocultos=True)] == ["oculto.zip"]


RAR = next(
    (r for n in ("Rar.exe", "rar") if (r := _buscar(n, _RUTAS_WINRAR))),
    None,
)


@pytest.mark.skipif(RAR is None, reason="se necesita Rar.exe (WinRAR) para crear un .rar")
def test_descomprime_rar(arbol: Path, tmp_path: Path) -> None:
    origen = tmp_path / "origen"
    origen.mkdir()
    crear_pdf(origen / "dentro.pdf", 4)
    subprocess.run(
        [RAR, "a", "-ep1", "-idq", str(arbol / "paquete.rar"), str(origen / "dentro.pdf")],
        check=True,
    )

    resultado = descomprimir_en_raiz(arbol)

    assert resultado.descomprimidos == ["paquete.rar"]
    assert (arbol / "paquete" / "dentro.pdf").is_file()


def test_rar_sin_herramientas_se_reporta(arbol: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("folder_reader.comprimidos._herramientas", lambda *_: [])
    (arbol / "paquete.rar").write_bytes(b"Rar!\x1a\x07\x00")

    resultado = descomprimir_en_raiz(arbol)

    assert resultado.fallidos[0][0] == "paquete.rar"
    assert "WinRAR o 7-Zip" in resultado.fallidos[0][1]


def _filas(destino: Path) -> list[dict[str, object]]:
    hoja = load_workbook(destino).active
    filas = list(hoja.iter_rows(values_only=True))
    return [dict(zip(filas[0], f)) for f in filas[1:]]


def test_la_cli_inventaria_el_contenido_descomprimido(arbol: Path, tmp_path: Path) -> None:
    contenido = tmp_path / "pdf_temporal.pdf"
    crear_pdf(contenido, 7)
    crear_zip(arbol / "anexos.zip", {"anexo.pdf": contenido.read_bytes()})

    destino = tmp_path / "salida.xlsx"
    assert cli.main([str(arbol), "-o", str(destino), "--sin-word"]) == 0

    por_nombre = {f["Nombre del archivo"]: f for f in _filas(destino)}
    assert por_nombre["anexo.pdf"]["Numero de paginas"] == 7
    assert por_nombre["anexo.pdf"]["Carpeta nivel 1"] == "anexos"
    # El .zip tambien sigue en el inventario como archivo.
    assert "anexos.zip" in por_nombre


def test_la_cli_puede_no_descomprimir(arbol: Path, tmp_path: Path) -> None:
    crear_zip(arbol / "anexos.zip", {"anexo.txt": b"a"})

    destino = tmp_path / "salida.xlsx"
    assert cli.main([str(arbol), "-o", str(destino), "--sin-descomprimir", "--sin-paginas"]) == 0

    assert not (arbol / "anexos").exists()
    assert "anexo.txt" not in {f["Nombre del archivo"] for f in _filas(destino)}
