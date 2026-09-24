from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from folder_reader.cli import main
from folder_reader.excel import COLUMNAS, Fila, escribir_excel


def leer(destino: Path) -> tuple[list[str], list[dict[str, object]]]:
    """Devuelve los encabezados y las filas como diccionarios.

    Acceder por nombre de columna evita que los tests se rompan cada vez que se
    inserta una columna nueva.
    """
    hoja = load_workbook(destino).active
    filas = list(hoja.iter_rows(values_only=True))
    encabezados = [str(c) for c in filas[0]]
    return encabezados, [dict(zip(encabezados, f)) for f in filas[1:]]


def fila_ejemplo(**cambios: object) -> Fila:
    base = dict(
        ruta=r"C:\docs\informe.pdf",
        ruta_relativa="informe.pdf",
        carpeta="docs",
        carpeta_nivel1="",
        # El caracter de control debe desaparecer al guardar.
        nombre="infor\x07me.pdf",
        extension=".pdf",
        tamano_bytes=2_097_152,
        creado=datetime(2025, 1, 2, 3, 4, 5),
        modificado=datetime(2025, 1, 2, 3, 4, 5),
        guardado=datetime(2024, 12, 31, 23, 59, 58),
        paginas=12,
        estado="OK",
    )
    return Fila(**{**base, **cambios})


def test_escribe_encabezado_y_tipos(tmp_path: Path) -> None:
    destino = tmp_path / "salida.xlsx"
    assert escribir_excel([fila_ejemplo()], destino) == 1

    encabezados, filas = leer(destino)
    assert encabezados == [titulo for titulo, _ in COLUMNAS]

    fila = filas[0]
    assert fila["Nombre del archivo"] == "informe.pdf"
    assert fila["Tamano (bytes)"] == 2_097_152
    assert fila["Tamano (MB)"] == 2.0
    assert isinstance(fila["Fecha de creacion"], datetime)
    assert fila["Guardado el (documento)"] == datetime(2024, 12, 31, 23, 59, 58)
    assert fila["Numero de paginas"] == 12

    hoja = load_workbook(destino).active
    assert hoja.freeze_panes == "A2"
    assert hoja.auto_filter.ref is not None


def test_las_tres_fechas_llevan_formato_de_fecha(tmp_path: Path) -> None:
    destino = tmp_path / "salida.xlsx"
    escribir_excel([fila_ejemplo()], destino)

    hoja = load_workbook(destino).active
    titulos = [c.value for c in hoja[1]]
    for titulo in ("Fecha de creacion", "Fecha de ultima modificacion", "Guardado el (documento)"):
        columna = titulos.index(titulo) + 1
        assert hoja.cell(row=2, column=columna).number_format == "yyyy-mm-dd hh:mm:ss"


def test_flujo_completo_desde_la_cli(arbol: Path, tmp_path: Path) -> None:
    destino = tmp_path / "inventario.xlsx"
    assert main([str(arbol), "-o", str(destino), "--sin-word"]) == 0

    _, filas = leer(destino)
    por_nombre = {f["Nombre del archivo"]: f for f in filas}

    assert "desktop.ini" not in por_nombre
    assert por_nombre["contrato.pdf"]["Numero de paginas"] == 3
    assert por_nombre["acta.docx"]["Numero de paginas"] == 5
    assert por_nombre["notas.txt"]["Estado"] == "No aplica"
    assert por_nombre["roto.pdf"]["Numero de paginas"] is None


def test_columnas_de_carpeta(arbol: Path, tmp_path: Path) -> None:
    import os

    destino = tmp_path / "inventario.xlsx"
    assert main([str(arbol), "-o", str(destino), "--sin-paginas"]) == 0

    _, filas = leer(destino)
    por_nombre = {f["Nombre del archivo"]: f for f in filas}

    contrato = por_nombre["contrato.pdf"]
    assert contrato["Carpeta contenedora"] == os.path.join("arbol", "contratos", "2025")
    assert contrato["Carpeta nivel 1"] == "contratos"

    anual = por_nombre["anual.pdf"]
    assert anual["Carpeta contenedora"] == os.path.join("arbol", "informes")
    assert anual["Carpeta nivel 1"] == "informes"

    # Un archivo en la propia raiz no tiene carpeta de primer nivel.
    assert por_nombre["portada.pdf"]["Carpeta contenedora"] == "arbol"
    assert por_nombre["portada.pdf"]["Carpeta nivel 1"] is None


def test_agrega_extension_xlsx_si_falta(arbol: Path, tmp_path: Path) -> None:
    destino = tmp_path / "inventario"
    assert main([str(arbol), "-o", str(destino), "--sin-paginas"]) == 0
    assert destino.with_suffix(".xlsx").exists()


def test_ruta_inexistente(tmp_path: Path) -> None:
    assert main([str(tmp_path / "no_existe")]) == 2
