from __future__ import annotations

from pathlib import Path

import pytest

from folder_reader import cli


def test_pregunta_hasta_recibir_una_carpeta_valida(
    arbol: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    respuestas = iter(["", "Z:/no_existe_en_ningun_lado", f'"{arbol}"'])
    monkeypatch.setattr("builtins.input", lambda *_: next(respuestas))

    assert cli._preguntar_ruta() == arbol

    errores = capsys.readouterr().err
    assert "Escriba una ruta" in errores
    assert "No existe o no es una carpeta" in errores


def test_salida_vacia_usa_el_valor_predeterminado(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("builtins.input", lambda *_: "  ")
    assert cli._preguntar_salida("inventario.xlsx") == "inventario.xlsx"


def test_normaliza_extensiones_sin_punto() -> None:
    assert cli._normalizar_extensiones(["PDF", ".Docx"]) == {".pdf", ".docx"}
    assert cli._normalizar_extensiones(None) is None


def test_nombre_predeterminado_lleva_la_carpeta(tmp_path: Path) -> None:
    nombre = cli._nombre_predeterminado(tmp_path / "Contratos")
    assert nombre.startswith("inventario_Contratos_")
    assert nombre.endswith(".xlsx")


def test_sin_ruta_y_sin_consola_devuelve_error(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    monkeypatch.setattr(cli.sys.stdin, "isatty", lambda: False)
    assert cli.main([]) == 2
    assert "Falta la carpeta" in capsys.readouterr().err


def test_carpeta_vacia_no_genera_excel(tmp_path: Path) -> None:
    vacia = tmp_path / "vacia"
    vacia.mkdir()
    destino = tmp_path / "salida.xlsx"
    assert cli.main([str(vacia), "-o", str(destino)]) == 1
    assert not destino.exists()


def test_filtro_de_extensiones_desde_la_cli(arbol: Path, tmp_path: Path) -> None:
    from openpyxl import load_workbook

    destino = tmp_path / "solo_pdf.xlsx"
    assert cli.main([str(arbol), "-o", str(destino), "--ext", "pdf"]) == 0

    hoja = load_workbook(destino).active
    titulos = [c.value for c in hoja[1]]
    columna = titulos.index("Extension")
    extensiones = {r[columna] for r in hoja.iter_rows(min_row=2, values_only=True)}
    assert extensiones == {".pdf"}
