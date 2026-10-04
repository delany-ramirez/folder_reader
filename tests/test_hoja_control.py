from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from folder_reader import cli_hoja_control
from folder_reader.hoja_control import (
    Documento,
    ErrorPlantilla,
    Grupo,
    clave_natural,
    escribir_hoja_control,
    plantilla_predeterminada,
)

# En la plantilla incluida el encabezado de columnas esta en la fila 5.
PRIMERA_FILA = 6


def _filas(destino: Path) -> list[tuple[object, ...]]:
    hoja = load_workbook(destino).active
    return [
        (fila[0].value, fila[1].value, fila[2].value, fila[8].value, fila[9].value)
        for fila in hoja.iter_rows(min_row=PRIMERA_FILA, max_col=10)
    ]


def test_orden_natural() -> None:
    nombres = ["10. CEDULA.pdf", "2. RP.pdf", "1. CONTRATO.pdf", "b.pdf", "A.pdf"]
    assert sorted(nombres, key=clave_natural) == [
        "1. CONTRATO.pdf",
        "2. RP.pdf",
        "10. CEDULA.pdf",
        "A.pdf",
        "b.pdf",
    ]


@pytest.mark.parametrize(
    ("desordenados", "esperado"),
    [
        (["1.jpeg", "10.jpeg", "11.jpeg", "2.jpeg", "3.jpeg"],
         ["1.jpeg", "2.jpeg", "3.jpeg", "10.jpeg", "11.jpeg"]),
        (["Informe 10", "Informe 2", "Informe 1"], ["Informe 1", "Informe 2", "Informe 10"]),
        (["2.10. Feria", "2.9. Charla", "2.1. Registro"],
         ["2.1. Registro", "2.9. Charla", "2.10. Feria"]),
        (["Foto (10).jpg", "Foto (2).jpg"], ["Foto (2).jpg", "Foto (10).jpg"]),
        (["１０. B", "２. A"], ["２. A", "１０. B"]),
    ],
)
def test_orden_numerico(desordenados: list[str], esperado: list[str]) -> None:
    assert sorted(desordenados, key=clave_natural) == esperado


def test_vuelca_carpetas_y_documentos_con_el_formato_de_la_plantilla(tmp_path: Path) -> None:
    destino = tmp_path / "hoja.xlsx"
    grupos = [
        Grupo("", [Documento("indice.pdf", date(2022, 1, 3), 1)]),
        Grupo(
            "1. PRECONTRACTUAL",
            [
                Documento("1. SOLICITUD.pdf", date(2022, 1, 12), 3),
                Documento("2. HOJA DE VIDA.pdf", date(2022, 1, 15), 6),
            ],
        ),
        Grupo("2. CONTRACTUAL", [Documento("CONTRATO.docx", None, None)]),
    ]

    assert escribir_hoja_control(grupos, destino) == 4

    hoja = load_workbook(destino).active
    filas = _filas(destino)
    assert filas[:6] == [
        (1, hoja["B6"].value, "indice.pdf", 1, hoja["J6"].value),
        (2, None, "1. PRECONTRACTUAL", None, None),
        (3, hoja["B8"].value, "1. SOLICITUD.pdf", 3, hoja["J8"].value),
        (4, hoja["B9"].value, "2. HOJA DE VIDA.pdf", 6, hoja["J9"].value),
        (5, None, "2. CONTRACTUAL", None, None),
        (6, None, "CONTRATO.docx", None, None),
    ]

    # Ambas fechas son la de guardado, sin hora y en dd/mm/yyyy.
    assert hoja["B8"].value.date() == date(2022, 1, 12) == hoja["J8"].value.date()
    assert hoja["B8"].number_format == hoja["J8"].number_format == "dd/mm/yyyy"

    # Solo las filas de carpeta van en negrilla.
    assert [hoja.cell(f, 3).font.b for f in range(6, 12)] == [
        False,
        True,
        False,
        False,
        True,
        False,
    ]

    # El total va justo debajo y suma el rango real; el pie de firmas lo sigue.
    assert hoja["C12"].value == "TOTAL"
    assert hoja["I12"].value == "=SUM(I6:I11)"
    assert hoja["A16"].value == "Diligenciado por: (Nombre completo)"

    combinadas = {str(r) for r in hoja.merged_cells.ranges}
    assert {f"C{f}:G{f}" for f in range(6, 13)} <= combinadas
    assert {"B1:H2", "B3:H4", "A16:D16", "G16:J16"} <= combinadas

    # No se agrega ninguna columna fuera del formato.
    assert hoja.max_column == 10

    # El encabezado (con el logo) se conserva intacto.
    assert hoja["B3"].value == "HOJA DE CONTROL"
    assert len(hoja._images) == 1


def test_no_quedan_restos_de_la_plantilla(tmp_path: Path) -> None:
    destino = tmp_path / "hoja.xlsx"
    escribir_hoja_control([Grupo("", [Documento("a.pdf", None, 2)])], destino)

    textos = [v for fila in _filas(destino) for v in fila if isinstance(v, str)]
    assert "NOMBRE DE LA CARPETA" not in textos
    assert "Nombre del documento.pdf" not in textos


def test_plantilla_sin_total_se_rechaza(tmp_path: Path) -> None:
    plantilla = tmp_path / "plantilla.xlsx"
    with plantilla_predeterminada().open("rb") as fh:
        libro = load_workbook(fh)
    libro.active["C8"].value = "SUMA"
    libro.save(plantilla)

    with pytest.raises(ErrorPlantilla, match="TOTAL"):
        escribir_hoja_control([], tmp_path / "hoja.xlsx", plantilla)


def test_cli_genera_un_archivo_por_carpeta_madre(
    arbol: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    salida = tmp_path / "salida"

    codigo = cli_hoja_control.main([str(arbol), "-o", str(salida), "--sin-word"])

    assert codigo == 0
    assert sorted(p.name for p in salida.iterdir()) == [
        "HOJA DE CONTROL contratos.xlsx",
        "HOJA DE CONTROL informes.xlsx",
    ]
    # contratos/2025/contrato.pdf: la subcarpeta 2025 sale como fila de carpeta.
    filas = _filas(salida / "HOJA DE CONTROL contratos.xlsx")
    assert [f[2] for f in filas[:3]] == ["2025", "contrato.pdf", "TOTAL"]
    assert filas[1][3] == 3

    # Los archivos de la propia carpeta madre van sin fila de carpeta.
    filas = _filas(salida / "HOJA DE CONTROL informes.xlsx")
    assert [f[2] for f in filas[:5]] == [
        "acta.docx",
        "anual.pdf",
        "borrador.docx",
        "roto.pdf",
        "TOTAL",
    ]
    # acta.docx trae su fecha de guardado en core.xml (18:30 UTC, mismo dia local).
    hoja = load_workbook(salida / "HOJA DE CONTROL informes.xlsx").active
    assert hoja["B6"].value.date() == date(2025, 3, 4)

    # anual.pdf no registra fecha de guardado: las dos fechas quedan en blanco.
    assert hoja["C7"].value == "anual.pdf"
    assert hoja["B7"].value is None and hoja["J7"].value is None
    assert hoja.max_column == 10

    salida_texto = capsys.readouterr().out
    assert "Archivos sueltos en la raiz, fuera de toda carpeta madre: 2" in salida_texto


def test_cli_con_carpeta_madre_genera_un_solo_archivo(arbol: Path, tmp_path: Path) -> None:
    salida = tmp_path / "salida"

    codigo = cli_hoja_control.main(
        [str(arbol), "-o", str(salida), "--sin-word", "--carpeta-madre"]
    )

    assert codigo == 0
    assert [p.name for p in salida.iterdir()] == ["HOJA DE CONTROL arbol.xlsx"]
    nombres = [f[2] for f in _filas(salida / "HOJA DE CONTROL arbol.xlsx")]
    assert nombres[:4] == ["notas.txt", "portada.pdf", "contratos\\2025", "contrato.pdf"]
